from __future__ import annotations

import io
import json
from datetime import UTC, date, datetime, timedelta
from typing import Annotated

from barcode import Code128
from barcode.writer import SVGWriter
from fastapi import APIRouter, Body, Depends, File, HTTPException, Query, UploadFile, status
from fastapi.responses import JSONResponse, StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.api.admin import Actor, require_any_permission, require_permission
from app.database import get_db
from app.models import (
    AccessRole,
    AuditLog,
    AuthSession,
    BusinessSetting,
    Cashier,
    CatalogOption,
    FinanceEntry,
    Member,
    Product,
    ProductBatch,
    Promotion,
    PurchaseOrder,
    Sale,
    SaleItem,
    SaleReturn,
    Shift,
    StockCount,
    Store,
)
from app.schemas import ProductCreate, ProductRead, ProductUpdate, PurchaseOrderRead, SaleRead
from app.schemas_operations import (
    AuditLogRead,
    AutoPurchaseOrderCreate,
    CashMovementCreate,
    CatalogOptionCreate,
    CatalogOptionRead,
    ChangePinCreate,
    FinanceEntryCreate,
    FinanceEntryRead,
    FinanceEntryUpdate,
    ForgotPasswordCreate,
    ForgotPasswordRead,
    LoginCreate,
    LoginRead,
    MemberAdminRead,
    MemberCreate,
    MemberUpdate,
    ProductBatchCreate,
    ProductBatchRead,
    ProfileUpdate,
    PromotionCreate,
    PromotionRead,
    PurchaseOrderDraftCreate,
    PurchaseReceiveCreate,
    PurchaseReturnCreate,
    ResetPasswordCreate,
    SettingRead,
    SettingUpdate,
    ShiftCloseCreate,
    ShiftOpenCreate,
    ShiftRead,
    StaffCreate,
    StockCountCreate,
    StockCountRead,
    StoreCreate,
    StoreProfileUpdate,
    TotpVerifyCreate,
    VoidSaleCreate,
)
from app.services.access import access_profile, permission_values
from app.services.operations import (
    add_cash_movement,
    authenticate,
    close_shift,
    create_batch,
    create_catalog_option,
    create_finance_entry,
    create_member,
    create_promotion,
    create_purchase_order,
    current_shift,
    hash_secret,
    now_utc,
    open_shift,
    perform_stock_count,
    receive_purchase_order,
    request_password_reset,
    reset_password,
    return_purchase,
    shift_summary,
    token_hash,
    update_member,
    verify_secret,
    verify_totp_challenge,
    void_sale,
    write_audit,
)
from app.services.products import create_product, update_product

router = APIRouter(tags=["Business operations"])
DbSession = Annotated[Session, Depends(get_db)]
PosActor = Annotated[Cashier, Depends(require_permission("pos.access"))]


def _ensure_own_cashier(actor: Cashier, cashier_code: str) -> None:
    if cashier_code.strip().upper() != actor.employee_code:
        raise HTTPException(
            status.HTTP_403_FORBIDDEN,
            "Identitas kasir tidak sesuai dengan sesi login",
        )


@router.post("/auth/login", response_model=LoginRead)
def login(payload: LoginCreate, session: DbSession) -> LoginRead:
    return authenticate(session, payload.identity, payload.password)


@router.post("/auth/totp/setup/confirm", response_model=LoginRead)
def confirm_totp_setup(payload: TotpVerifyCreate, session: DbSession) -> LoginRead:
    return verify_totp_challenge(
        session,
        payload.challenge_token,
        payload.code,
        expected_purpose="setup",
    )


@router.post("/auth/totp/verify", response_model=LoginRead)
def verify_totp_login(payload: TotpVerifyCreate, session: DbSession) -> LoginRead:
    return verify_totp_challenge(
        session,
        payload.challenge_token,
        payload.code,
        expected_purpose="login",
    )


@router.get("/products/{product_id}/barcode.svg")
def product_barcode(product_id: str, session: DbSession) -> StreamingResponse:
    product = session.get(Product, product_id)
    if product is None or not product.is_active:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Produk tidak ditemukan")
    output = io.BytesIO()
    Code128(product.barcode, writer=SVGWriter()).write(
        output,
        options={
            "module_width": 0.25,
            "module_height": 12,
            "font_size": 0,
            "text_distance": 0,
            "quiet_zone": 2,
        },
    )
    output.seek(0)
    return StreamingResponse(output, media_type="image/svg+xml")


@router.get("/pos/config")
def pos_config(session: DbSession, _actor: PosActor) -> dict[str, object]:
    store = session.scalar(select(Store).where(Store.code == "MY-001"))
    settings = {item.key: item.value for item in session.scalars(select(BusinessSetting)).all()}
    methods = [
        item.strip()
        for item in settings.get("payment.methods", "cash,qris,edc").split(",")
        if item.strip()
    ]
    try:
        tax_rate = max(0, int(settings.get("tax.rate", "0")))
    except ValueError:
        tax_rate = 0
    return {
        "store_name": store.name if store else "MYMARKET",
        "store_address": store.address if store else None,
        "logo_url": store.logo_url if store and store.logo_url else "/logo-app.png",
        "payment_methods": methods,
        "tax_rate": tax_rate,
        "tax_included": settings.get("tax.included", "true").lower() == "true",
        "receipt_paper": settings.get("receipt.paper", "80"),
        "receipt_footer": settings.get("receipt.footer", "Terima kasih"),
    }


@router.post("/auth/forgot-password", response_model=ForgotPasswordRead)
def forgot_password(payload: ForgotPasswordCreate, session: DbSession) -> ForgotPasswordRead:
    token = request_password_reset(session, payload.email)
    return ForgotPasswordRead(
        message="Jika email terdaftar, link reset sudah dibuat.",
        reset_token=token,
    )


@router.post("/auth/reset-password", status_code=status.HTTP_204_NO_CONTENT)
def apply_password_reset(payload: ResetPasswordCreate, session: DbSession) -> None:
    reset_password(session, payload.token, payload.new_password)


@router.get("/auth/profile")
def profile(actor: Actor) -> dict[str, object]:
    result = access_profile(actor).model_dump()
    result["email"] = actor.email
    return result


@router.put("/auth/profile")
def update_profile(payload: ProfileUpdate, session: DbSession, actor: Actor) -> dict[str, object]:
    duplicate = session.scalar(
        select(Cashier).where(Cashier.email == payload.email.lower(), Cashier.id != actor.id)
    )
    if duplicate:
        raise HTTPException(status.HTTP_409_CONFLICT, "Email sudah digunakan")
    actor.full_name = payload.full_name.strip()
    actor.email = payload.email.lower()
    write_audit(session, actor, "account.profile", "cashier", actor.id)
    session.commit()
    return profile(actor)


@router.post("/auth/change-pin", status_code=status.HTTP_204_NO_CONTENT)
def change_pin(payload: ChangePinCreate, session: DbSession, actor: Actor) -> None:
    if not verify_secret(payload.current_pin, actor.pin_hash):
        raise HTTPException(status.HTTP_403_FORBIDDEN, "PIN lama salah")
    actor.pin_hash = hash_secret(payload.new_pin)
    write_audit(session, actor, "account.pin", "cashier", actor.id)
    session.commit()


@router.post("/auth/logout", status_code=status.HTTP_204_NO_CONTENT)
def logout(
    session: DbSession,
    token: Annotated[str | None, Body(embed=True)] = None,
) -> None:
    if not token:
        return
    auth_session = session.scalar(
        select(AuthSession).where(
            AuthSession.token_hash == token_hash(token),
            AuthSession.revoked_at.is_(None),
        )
    )
    if auth_session:
        auth_session.revoked_at = now_utc()
        session.commit()


@router.get("/shifts/current", response_model=ShiftRead | None)
def get_current_shift(
    session: DbSession,
    actor: PosActor,
    cashier_code: Annotated[str, Query(max_length=30)],
    store_code: Annotated[str, Query(max_length=20)] = "MY-001",
) -> ShiftRead | None:
    _ensure_own_cashier(actor, cashier_code)
    return current_shift(session, store_code, cashier_code)


@router.post("/shifts/open", response_model=ShiftRead, status_code=status.HTTP_201_CREATED)
def create_shift(payload: ShiftOpenCreate, session: DbSession, actor: PosActor) -> ShiftRead:
    _ensure_own_cashier(actor, payload.cashier_code)
    return open_shift(session, payload.store_code, payload.cashier_code, payload.opening_cash)


@router.post("/shifts/{shift_id}/cash-movements", response_model=ShiftRead)
def create_cash_movement(
    shift_id: str, payload: CashMovementCreate, session: DbSession, actor: PosActor
) -> ShiftRead:
    _ensure_own_cashier(actor, payload.cashier_code)
    return add_cash_movement(session, shift_id, payload, payload.cashier_code)


@router.post("/shifts/{shift_id}/close", response_model=ShiftRead)
def finish_shift(
    shift_id: str, payload: ShiftCloseCreate, session: DbSession, actor: PosActor
) -> ShiftRead:
    _ensure_own_cashier(actor, payload.cashier_code)
    return close_shift(session, shift_id, payload.cashier_code, payload.closing_cash)


@router.get("/cashier/sales", response_model=list[SaleRead])
def cashier_sales(
    session: DbSession,
    actor: PosActor,
    cashier_code: Annotated[str, Query(max_length=30)],
    limit: Annotated[int, Query(ge=1, le=100)] = 30,
) -> list[Sale]:
    _ensure_own_cashier(actor, cashier_code)
    cashier = session.scalar(select(Cashier).where(Cashier.employee_code == cashier_code))
    if cashier is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Kasir tidak ditemukan")
    statement = (
        select(Sale)
        .where(Sale.cashier_id == cashier.id, Sale.status.in_(("completed", "cancelled")))
        .order_by(Sale.created_at.desc())
        .limit(limit)
    )
    return list(session.scalars(statement).all())


@router.delete("/sales/holds/{sale_id}", status_code=status.HTTP_204_NO_CONTENT)
def remove_hold(sale_id: str, session: DbSession, actor: PosActor) -> None:
    sale = session.get(Sale, sale_id)
    if sale is None or sale.status != "held" or sale.cashier_id != actor.id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Transaksi hold tidak ditemukan")
    session.delete(sale)
    session.commit()


@router.post("/sales/{transaction_no}/void")
def cancel_sale(
    transaction_no: str, payload: VoidSaleCreate, session: DbSession, _actor: PosActor
) -> dict[str, object]:
    item = void_sale(
        session,
        transaction_no,
        payload.supervisor_code,
        payload.supervisor_pin,
        payload.reason,
    )
    return {"id": item.id, "return_no": item.return_no, "amount": item.amount}


@router.get("/admin/operations")
def operations_snapshot(
    session: DbSession,
    actor: Annotated[Cashier, Depends(require_any_permission("admin.operations", "admin.finance"))],
) -> dict[str, object]:
    today = datetime.now(UTC).date()
    permissions = permission_values(actor.access_role)
    can_view_finance = "admin.finance" in permissions
    can_view_settings = "admin.settings" in permissions
    batches = session.execute(
        select(ProductBatch, Product)
        .join(Product, Product.id == ProductBatch.product_id)
        .order_by(ProductBatch.expiry_date)
        .limit(100)
    ).all()
    audits = session.execute(
        select(AuditLog, Cashier.full_name)
        .outerjoin(Cashier, Cashier.id == AuditLog.actor_id)
        .order_by(AuditLog.created_at.desc())
        .limit(100)
    ).all()
    shifts = list(session.scalars(select(Shift).order_by(Shift.opened_at.desc()).limit(50)).all())
    return {
        "stores": [
            {
                "id": item.id,
                "code": item.code,
                "name": item.name,
                "address": item.address,
                "timezone": item.timezone,
                "logo_url": item.logo_url,
                "tax_id": item.tax_id,
            }
            for item in session.scalars(select(Store).order_by(Store.code)).all()
        ],
        "catalog_options": [
            CatalogOptionRead.model_validate(item).model_dump(mode="json")
            for item in session.scalars(
                select(CatalogOption).order_by(CatalogOption.kind, CatalogOption.name)
            ).all()
        ],
        "promotions": [
            PromotionRead.model_validate(item).model_dump(mode="json")
            for item in session.scalars(
                select(Promotion).order_by(Promotion.created_at.desc())
            ).all()
        ],
        "batches": [
            ProductBatchRead(
                id=batch.id,
                product_id=batch.product_id,
                product_name=product.name,
                barcode=product.barcode,
                batch_no=batch.batch_no,
                expiry_date=batch.expiry_date,
                quantity=batch.quantity,
                days_until_expiry=(batch.expiry_date - today).days,
                status="expired"
                if batch.expiry_date < today
                else "warning"
                if batch.expiry_date <= today + timedelta(days=30)
                else "safe",
            ).model_dump(mode="json")
            for batch, product in batches
        ],
        "stock_counts": [
            StockCountRead.model_validate(item).model_dump(mode="json")
            for item in session.scalars(
                select(StockCount).order_by(StockCount.created_at.desc()).limit(100)
            ).all()
        ],
        "members": [
            MemberAdminRead.model_validate(item).model_dump(mode="json")
            for item in session.scalars(
                select(Member).order_by(Member.created_at.desc()).limit(200)
            ).all()
        ],
        "finance_entries": [
            FinanceEntryRead.model_validate(item).model_dump(mode="json")
            for item in session.scalars(
                select(FinanceEntry).order_by(FinanceEntry.created_at.desc()).limit(200)
            ).all()
        ]
        if can_view_finance
        else [],
        "settings": [
            SettingRead.model_validate(item).model_dump(mode="json")
            for item in session.scalars(
                select(BusinessSetting).order_by(BusinessSetting.group_name, BusinessSetting.key)
            ).all()
        ]
        if can_view_settings
        else [],
        "shifts": [shift_summary(session, shift).model_dump(mode="json") for shift in shifts],
        "returns": [
            {
                "id": item.id,
                "return_no": item.return_no,
                "sale_id": item.sale_id,
                "amount": item.amount,
                "reason": item.reason,
                "created_at": item.created_at.isoformat(),
            }
            for item in session.scalars(
                select(SaleReturn).order_by(SaleReturn.created_at.desc()).limit(100)
            ).all()
        ],
        "audit_logs": [
            AuditLogRead(
                id=log.id,
                actor_name=actor_name,
                action=log.action,
                entity_type=log.entity_type,
                entity_id=log.entity_id,
                details=json.loads(log.details_json or "{}"),
                created_at=log.created_at,
            ).model_dump(mode="json")
            for log, actor_name in audits
        ]
        if can_view_settings
        else [],
    }


@router.post("/admin/catalog-options", response_model=CatalogOptionRead)
def add_catalog_option(
    payload: CatalogOptionCreate,
    session: DbSession,
    actor: Annotated[Cashier, Depends(require_permission("admin.products"))],
) -> CatalogOption:
    return create_catalog_option(session, payload, actor)


@router.post("/admin/promotions", response_model=PromotionRead)
def add_promotion(
    payload: PromotionCreate,
    session: DbSession,
    actor: Annotated[Cashier, Depends(require_permission("admin.products"))],
) -> Promotion:
    return create_promotion(session, payload, actor)


@router.post("/admin/product-batches", response_model=ProductBatchRead)
def add_product_batch(
    payload: ProductBatchCreate,
    session: DbSession,
    actor: Annotated[Cashier, Depends(require_permission("admin.stock"))],
) -> ProductBatchRead:
    batch = create_batch(session, payload, actor)
    product = session.get(Product, batch.product_id)
    days = (batch.expiry_date - datetime.now(UTC).date()).days
    return ProductBatchRead(
        id=batch.id,
        product_id=batch.product_id,
        product_name=product.name if product else "-",
        barcode=product.barcode if product else "-",
        batch_no=batch.batch_no,
        expiry_date=batch.expiry_date,
        quantity=batch.quantity,
        days_until_expiry=days,
        status="expired" if days < 0 else "warning" if days <= 30 else "safe",
    )


@router.post("/admin/stock-counts", response_model=StockCountRead)
def add_stock_count(
    payload: StockCountCreate,
    session: DbSession,
    actor: Annotated[Cashier, Depends(require_permission("admin.stock"))],
) -> StockCount:
    return perform_stock_count(session, payload, actor)


@router.post("/admin/members", response_model=MemberAdminRead)
def add_member(
    payload: MemberCreate,
    session: DbSession,
    actor: Annotated[Cashier, Depends(require_permission("admin.operations"))],
) -> Member:
    return create_member(session, payload, actor)


@router.put("/admin/members/{member_id}", response_model=MemberAdminRead)
def edit_member(
    member_id: str,
    payload: MemberUpdate,
    session: DbSession,
    actor: Annotated[Cashier, Depends(require_permission("admin.operations"))],
) -> Member:
    member = session.get(Member, member_id)
    if member is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Member tidak ditemukan")
    return update_member(session, member, payload, actor)


@router.post("/admin/finance", response_model=FinanceEntryRead)
def add_finance_entry(
    payload: FinanceEntryCreate,
    session: DbSession,
    actor: Annotated[Cashier, Depends(require_permission("admin.finance"))],
) -> FinanceEntry:
    return create_finance_entry(session, payload, actor)


@router.put("/admin/finance/{entry_id}", response_model=FinanceEntryRead)
def settle_finance_entry(
    entry_id: str,
    payload: FinanceEntryUpdate,
    session: DbSession,
    actor: Annotated[Cashier, Depends(require_permission("admin.finance"))],
) -> FinanceEntry:
    entry = session.get(FinanceEntry, entry_id)
    if entry is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Catatan keuangan tidak ditemukan")
    entry.status = payload.status
    write_audit(
        session, actor, "finance.status", "finance_entry", entry.id, {"status": payload.status}
    )
    session.commit()
    session.refresh(entry)
    return entry


@router.post("/admin/purchase-orders", response_model=PurchaseOrderRead)
def add_purchase_order(
    payload: PurchaseOrderDraftCreate,
    session: DbSession,
    actor: Annotated[Cashier, Depends(require_permission("admin.purchasing"))],
) -> PurchaseOrder:
    return create_purchase_order(session, payload, actor)


@router.post("/admin/purchase-orders/auto", response_model=PurchaseOrderRead)
def create_automatic_purchase_order(
    payload: AutoPurchaseOrderCreate,
    session: DbSession,
    actor: Annotated[Cashier, Depends(require_permission("admin.purchasing"))],
) -> PurchaseOrder:
    products = session.scalars(
        select(Product).where(
            Product.is_active.is_(True),
            Product.stock + Product.warehouse_stock <= Product.reorder_level,
        )
    ).all()
    if not products:
        raise HTTPException(status.HTTP_409_CONFLICT, "Semua stok masih di atas batas minimum")
    suggested = PurchaseOrderDraftCreate(
        supplier_id=payload.supplier_id,
        store_code=payload.store_code,
        notes="PO otomatis dari stok minimum",
        items=[
            {
                "barcode": product.barcode,
                "quantity": max(
                    product.reorder_level * 2 - product.stock - product.warehouse_stock, 1
                ),
                "unit_cost": product.cost_price,
            }
            for product in products
        ],
    )
    return create_purchase_order(session, suggested, actor)


@router.post("/admin/purchase-orders/{order_id}/receive", response_model=PurchaseOrderRead)
def receive_order(
    order_id: str,
    payload: PurchaseReceiveCreate,
    session: DbSession,
    actor: Annotated[Cashier, Depends(require_permission("admin.purchasing"))],
) -> PurchaseOrder:
    order = session.get(PurchaseOrder, order_id)
    if order is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Purchase order tidak ditemukan")
    return receive_purchase_order(session, order, payload, actor)


@router.post("/admin/purchase-returns", response_model=FinanceEntryRead)
def add_purchase_return(
    payload: PurchaseReturnCreate,
    session: DbSession,
    actor: Annotated[Cashier, Depends(require_permission("admin.purchasing"))],
) -> FinanceEntry:
    return return_purchase(session, payload, actor)


@router.put("/admin/settings/{key}", response_model=SettingRead)
def save_setting(
    key: str,
    payload: SettingUpdate,
    session: DbSession,
    actor: Annotated[Cashier, Depends(require_permission("admin.settings"))],
) -> BusinessSetting:
    setting = session.get(BusinessSetting, key)
    if setting is None:
        setting = BusinessSetting(key=key)
        session.add(setting)
    setting.group_name = payload.group_name
    setting.label = payload.label
    setting.value = payload.value
    write_audit(session, actor, "setting.update", "business_setting", key)
    session.commit()
    session.refresh(setting)
    return setting


@router.put("/admin/store-profile")
def save_store_profile(
    payload: StoreProfileUpdate,
    session: DbSession,
    actor: Annotated[Cashier, Depends(require_permission("admin.settings"))],
) -> dict[str, object]:
    store = session.scalar(select(Store).where(Store.code == "MY-001"))
    if store is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Toko tidak ditemukan")
    for field, value in payload.model_dump().items():
        setattr(store, field, value)
    write_audit(session, actor, "store.profile", "store", store.id)
    session.commit()
    return {"code": store.code, **payload.model_dump()}


@router.post("/admin/stores", status_code=status.HTTP_201_CREATED)
def add_store(
    payload: StoreCreate,
    session: DbSession,
    actor: Annotated[Cashier, Depends(require_permission("admin.settings"))],
) -> dict[str, object]:
    code = payload.code.strip().upper()
    if session.scalar(select(Store).where(Store.code == code)):
        raise HTTPException(status.HTTP_409_CONFLICT, "Kode cabang sudah digunakan")
    values = payload.model_dump(exclude={"code"})
    store = Store(code=code, **values)
    session.add(store)
    session.flush()
    write_audit(session, actor, "store.create", "store", store.id, {"code": code})
    session.commit()
    return {"id": store.id, "code": store.code, "name": store.name}


@router.post("/admin/staff")
def add_staff(
    payload: StaffCreate,
    session: DbSession,
    actor: Annotated[Cashier, Depends(require_permission("admin.roles"))],
) -> dict[str, object]:
    code = payload.employee_code.strip().upper()
    email = payload.email.strip().lower()
    if session.scalar(
        select(Cashier).where((Cashier.employee_code == code) | (Cashier.email == email))
    ):
        raise HTTPException(status.HTTP_409_CONFLICT, "Kode pegawai atau email sudah digunakan")
    role = session.get(AccessRole, payload.role_id)
    if role is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Role tidak ditemukan")
    employee = Cashier(
        employee_code=code,
        full_name=payload.full_name.strip(),
        email=email,
        password_hash=hash_secret(payload.password),
        pin_hash=hash_secret(payload.pin),
        role=role.code.lower(),
        role_id=role.id,
        is_active=True,
    )
    session.add(employee)
    session.flush()
    write_audit(session, actor, "staff.create", "cashier", employee.id, {"code": code})
    session.commit()
    return {
        "id": employee.id,
        "employee_code": code,
        "full_name": employee.full_name,
        "email": email,
        "role_name": role.name,
    }


@router.get("/admin/reports/operations")
def operations_report(
    session: DbSession,
    _actor: Annotated[Cashier, Depends(require_permission("admin.reports"))],
) -> dict[str, object]:
    start = datetime.combine(datetime.now(UTC).date().replace(day=1), datetime.min.time())
    revenue = int(
        session.scalar(
            select(func.coalesce(func.sum(Sale.total), 0)).where(
                Sale.status == "completed", Sale.created_at >= start
            )
        )
        or 0
    )
    cost = int(
        session.scalar(
            select(func.coalesce(func.sum(SaleItem.quantity * Product.cost_price), 0))
            .join(Sale, Sale.id == SaleItem.sale_id)
            .join(Product, Product.id == SaleItem.product_id)
            .where(Sale.status == "completed", Sale.created_at >= start)
        )
        or 0
    )
    expenses = int(
        session.scalar(
            select(func.coalesce(func.sum(FinanceEntry.amount), 0)).where(
                FinanceEntry.entry_type == "expense",
                FinanceEntry.status != "cancelled",
                FinanceEntry.created_at >= start,
            )
        )
        or 0
    )
    inventory_value = int(
        session.scalar(
            select(
                func.coalesce(
                    func.sum((Product.stock + Product.warehouse_stock) * Product.cost_price), 0
                )
            )
        )
        or 0
    )
    sold = (
        select(SaleItem.product_id, func.coalesce(func.sum(SaleItem.quantity), 0).label("sold"))
        .join(Sale, Sale.id == SaleItem.sale_id)
        .where(Sale.status == "completed", Sale.created_at >= start)
        .group_by(SaleItem.product_id)
        .subquery()
    )
    slow = session.execute(
        select(Product.name, Product.stock, Product.warehouse_stock, func.coalesce(sold.c.sold, 0))
        .outerjoin(sold, sold.c.product_id == Product.id)
        .where(Product.is_active.is_(True))
        .order_by(func.coalesce(sold.c.sold, 0), Product.name)
        .limit(10)
    ).all()
    tax_rate_setting = session.get(BusinessSetting, "tax.rate")
    tax_included_setting = session.get(BusinessSetting, "tax.included")
    try:
        tax_rate = max(0, int(tax_rate_setting.value if tax_rate_setting else "0"))
    except ValueError:
        tax_rate = 0
    tax_included = (
        tax_included_setting.value if tax_included_setting else "true"
    ).lower() == "true"
    tax_estimate = (
        revenue * tax_rate // (100 + tax_rate)
        if tax_included and tax_rate
        else revenue * tax_rate // 100
    )
    purchase_spend = int(
        session.scalar(
            select(func.coalesce(func.sum(PurchaseOrder.subtotal), 0)).where(
                PurchaseOrder.status.in_(("partially_received", "received")),
                PurchaseOrder.created_at >= start,
            )
        )
        or 0
    )
    cashier_sales = session.execute(
        select(
            Cashier.full_name,
            func.count(Sale.id),
            func.coalesce(func.sum(Sale.total), 0),
        )
        .join(Sale, Sale.cashier_id == Cashier.id)
        .where(Sale.status == "completed", Sale.created_at >= start)
        .group_by(Cashier.id, Cashier.full_name)
        .order_by(func.sum(Sale.total).desc())
    ).all()
    return {
        "period": start.date().isoformat(),
        "revenue": revenue,
        "gross_profit": revenue - cost,
        "expenses": expenses,
        "net_profit": revenue - cost - expenses,
        "inventory_value": inventory_value,
        "tax_estimate": tax_estimate,
        "purchase_spend": purchase_spend,
        "cashiers": [
            {"name": name, "transactions": int(transactions), "revenue": int(total)}
            for name, transactions, total in cashier_sales
        ],
        "open_payables": int(
            session.scalar(
                select(func.coalesce(func.sum(FinanceEntry.amount), 0)).where(
                    FinanceEntry.entry_type == "payable", FinanceEntry.status == "open"
                )
            )
            or 0
        ),
        "open_receivables": int(
            session.scalar(
                select(func.coalesce(func.sum(FinanceEntry.amount), 0)).where(
                    FinanceEntry.entry_type == "receivable", FinanceEntry.status == "open"
                )
            )
            or 0
        ),
        "slow_moving": [
            {"name": name, "stock": stock + warehouse, "sold": int(quantity)}
            for name, stock, warehouse, quantity in slow
        ],
    }


@router.get("/admin/products/export.xlsx")
def export_products(
    session: DbSession,
    _actor: Annotated[Cashier, Depends(require_permission("admin.products"))],
) -> StreamingResponse:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Produk"
    sheet.append(
        [
            "barcode",
            "sku",
            "nama",
            "kategori",
            "merek",
            "satuan",
            "konversi",
            "lokasi",
            "harga",
            "modal",
            "stok",
            "stok_gudang",
            "stok_minimum",
            "aktif",
        ]
    )
    for product in session.scalars(
        select(Product).where(Product.is_active.is_(True)).order_by(Product.name)
    ).all():
        sheet.append(
            [
                product.barcode,
                product.sku,
                product.name,
                product.category,
                product.brand,
                product.unit,
                product.unit_conversion,
                product.location,
                product.price,
                product.cost_price,
                product.stock,
                product.warehouse_stock,
                product.reorder_level,
                product.is_active,
            ]
        )
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=produk-mymarket.xlsx"},
    )


@router.post("/admin/products/import.xlsx")
def import_products(
    session: DbSession,
    actor: Annotated[Cashier, Depends(require_permission("admin.products"))],
    file: Annotated[UploadFile, File()],
) -> dict[str, int]:
    try:
        workbook = load_workbook(file.file, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
    except Exception as exc:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_CONTENT, "File Excel tidak valid"
        ) from exc
    if not rows:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, "File Excel kosong")
    headers = [str(value or "").strip().lower() for value in rows[0]]
    required = {"barcode", "nama", "kategori", "lokasi", "harga", "modal", "stok_minimum"}
    if not required.issubset(headers):
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_CONTENT, "Kolom wajib Excel belum lengkap"
        )
    created = updated = 0
    for raw_row in rows[1:]:
        row = dict(zip(headers, raw_row, strict=False))
        if not row.get("barcode"):
            continue
        existing = session.scalar(select(Product).where(Product.barcode == str(row["barcode"])))
        values = {
            "barcode": str(row["barcode"]),
            "sku": str(row.get("sku") or "") or None,
            "name": str(row.get("nama") or "").strip(),
            "category": str(row.get("kategori") or "Umum").strip(),
            "brand": str(row.get("merek") or "").strip() or None,
            "unit": str(row.get("satuan") or "pcs").strip(),
            "unit_conversion": int(row.get("konversi") or 1),
            "location": str(row.get("lokasi") or "Gudang").strip(),
            "price": int(row.get("harga") or 0),
            "cost_price": int(row.get("modal") or 0),
            "reorder_level": int(row.get("stok_minimum") or 0),
            "is_active": bool(True if row.get("aktif") is None else row.get("aktif")),
        }
        if existing:
            update_product(session, existing, ProductUpdate(**values))
            updated += 1
        else:
            create_product(session, ProductCreate(**values))
            created += 1
    write_audit(
        session, actor, "product.import", "product", None, {"created": created, "updated": updated}
    )
    session.commit()
    return {"created": created, "updated": updated}


@router.get("/admin/backup")
def backup_data(
    session: DbSession,
    _actor: Annotated[Cashier, Depends(require_permission("admin.settings"))],
) -> JSONResponse:
    payload = {
        "generated_at": datetime.now(UTC).isoformat(),
        "stores": [
            {"code": item.code, "name": item.name, "address": item.address}
            for item in session.scalars(select(Store)).all()
        ],
        "products": [
            ProductRead.model_validate(item).model_dump(mode="json")
            for item in session.scalars(select(Product)).all()
        ],
        "members": [
            MemberAdminRead.model_validate(item).model_dump(mode="json")
            for item in session.scalars(select(Member)).all()
        ],
        "sales": [
            {
                "transaction_no": item.transaction_no,
                "status": item.status,
                "total": item.total,
                "created_at": item.created_at.isoformat(),
            }
            for item in session.scalars(select(Sale)).all()
        ],
        "finance": [
            FinanceEntryRead.model_validate(item).model_dump(mode="json")
            for item in session.scalars(select(FinanceEntry)).all()
        ],
    }
    response = JSONResponse(payload)
    response.headers["Content-Disposition"] = (
        f"attachment; filename=mymarket-backup-{date.today():%Y%m%d}.json"
    )
    return response
